#!/usr/bin/env python3
"""`data/training/raw/`의 원본 hwp/xlsx 파일을 AI가 읽기 쉬운 평문 텍스트/CSV로 변환한다.

- .hwp: `hwp5proc xml`로 전체 구조를 XML로 뽑은 뒤, 문단·표(표 안 문단 포함)를
  순서대로 훑어 하나의 텍스트 파일로 만든다. 표 안 내용도 `<표>` 같은 자리표시자로
  버리지 않고 행/셀 텍스트를 그대로 옮긴다(`pyhwp`의 `hwp5txt` 명령은 표 내용을
  버리므로 쓰지 않았다).
- .xlsx: 시트별로 모든 셀 값을 그대로 CSV로 저장한다(수식이 아니라 계산된 값 기준).

변환 후 원본 XML의 <Text> 요소에 들어있는 한글 음절 수와 변환 결과물의 한글 음절
수를 세어 비교하고, 일치하지 않으면 경고를 띄운다 — 표 등 어딘가에서 텍스트가
빠졌는지 확인하는 손실 검증용이다.

출력은 `data/training/raw/converted/`에 쌓이며, 원본과 마찬가지로 개인정보라
`.gitignore`(`data/training/raw/`)로 보호된다.

사용법:
    python tools/convert_raw_training_data.py
"""
from __future__ import annotations

import csv
import subprocess
import sys
from pathlib import Path

from lxml import etree
import openpyxl

RAW_DIR = Path(__file__).resolve().parent.parent / 'data' / 'training' / 'raw'
OUT_DIR = RAW_DIR / 'converted'

HANGUL_RANGE = (0xAC00, 0xD7A3)


def count_hangul(text: str) -> int:
    return sum(1 for ch in text if HANGUL_RANGE[0] <= ord(ch) <= HANGUL_RANGE[1])


# ---- HWP -------------------------------------------------------------

def render_node(el: etree._Element) -> str:
    tag = el.tag
    if tag == 'Text':
        return el.text or ''
    if tag == 'ControlChar':
        code = el.get('code')
        if code == '13':  # PARAGRAPH_BREAK
            return '\n'
        if code == '10':  # LINE_BREAK
            return '\n'
        if code == '9':  # TAB
            return '\t'
        return ''
    if tag == 'TableRow':
        cells = [render_node(c).strip().replace('\n', ' ') for c in el.findall('TableCell')]
        return ' | '.join(cells) + '\n'
    if tag == 'TableControl':
        tbody = el.find('TableBody')
        if tbody is None:
            return ''
        rows = ''.join(render_node(r) for r in tbody.findall('TableRow'))
        return '\n<TABLE>\n' + rows + '</TABLE>\n'
    return ''.join(render_node(c) for c in el)


def hwp_to_text(path: Path) -> tuple[str, int]:
    """(변환된 텍스트, 원본 XML의 <Text> 요소 한글 음절 수)를 반환한다."""
    xml_bytes = subprocess.run(
        ['hwp5proc', 'xml', str(path)], capture_output=True, check=True,
    ).stdout
    root = etree.fromstring(xml_bytes)
    body = root.find('.//BodyText')
    if body is None:
        raise RuntimeError(f'{path.name}: BodyText를 찾을 수 없음')

    source_hangul = sum(count_hangul(t.text or '') for t in body.iter('Text'))

    sections = body.findall('SectionDef')
    text = '\n\n'.join(render_node(s) for s in sections)
    # 문단 구분이 과하게 벌어지는 걸 정리
    while '\n\n\n' in text:
        text = text.replace('\n\n\n', '\n\n')
    return text.strip() + '\n', source_hangul


def convert_hwp_files() -> None:
    for path in sorted(RAW_DIR.glob('*.hwp')):
        try:
            text, source_hangul = hwp_to_text(path)
        except (subprocess.CalledProcessError, FileNotFoundError, RuntimeError) as exc:
            # FileNotFoundError는 `hwp5proc` 실행파일 자체가 없을 때(PATH 미설정 등) —
            # 이 경우 이후 파일도 전부 같은 이유로 실패하지만, 그래도 xlsx 변환은
            # 계속 진행되도록 여기서 계속(continue)한다.
            print(f'[hwp] {path.name} -> 변환 실패: {exc}')
            continue
        out_path = OUT_DIR / f'{path.stem}.txt'
        out_path.write_text(text, encoding='utf-8')

        out_hangul = count_hangul(text)
        status = 'OK' if out_hangul == source_hangul else '불일치!'
        print(f'[hwp] {path.name} -> {out_path.name} '
              f'(한글 음절: 원본 {source_hangul} / 변환 {out_hangul} {status})')


# ---- XLSX --------------------------------------------------------------

def convert_xlsx_files() -> None:
    for path in sorted(RAW_DIR.glob('*.xlsx')):
        try:
            wb = openpyxl.load_workbook(path, data_only=True)
        except Exception as exc:  # noqa: BLE001 — 손상/암호화된 파일 하나 때문에 나머지가 안 돌면 안 됨
            print(f'[xlsx] {path.name} -> 변환 실패: {exc}')
            continue
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            safe_sheet = sheet_name.replace('/', '_')
            out_path = OUT_DIR / f'{path.stem}__{safe_sheet}.csv'

            source_hangul = 0
            rows_written = 0
            with out_path.open('w', encoding='utf-8-sig', newline='') as f:
                writer = csv.writer(f)
                for row in ws.iter_rows():
                    values = ['' if c.value is None else str(c.value) for c in row]
                    if not any(v.strip() for v in values):
                        continue
                    writer.writerow(values)
                    rows_written += 1
                    for v in values:
                        source_hangul += count_hangul(v)

            out_text = out_path.read_text(encoding='utf-8-sig')
            out_hangul = count_hangul(out_text)
            status = 'OK' if out_hangul == source_hangul else '불일치!'
            print(f'[xlsx] {path.name} [{sheet_name}] -> {out_path.name} '
                  f'({rows_written}행, 한글 음절: 원본 {source_hangul} / 변환 {out_hangul} {status})')


def main() -> None:
    if not RAW_DIR.exists():
        print(f'{RAW_DIR}가 없습니다.')
        sys.exit(1)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    convert_hwp_files()
    convert_xlsx_files()


if __name__ == '__main__':
    main()
