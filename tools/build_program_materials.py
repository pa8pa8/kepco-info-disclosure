from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT.parent / 'INFO_DISCLOSURE_System_materials'
DOCX_PATH = OUT / 'INFO_DISCLOSURE_System_Security_Review.docx'


def doc_set_cell_text(cell, text, bold=False, color='111827', size=9):
    cell.text = ''
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.bold = bold
    r.font.name = '맑은 고딕'
    r._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)


def doc_set_cell_shading(cell, fill):
    from docx.oxml import OxmlElement
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def build_security_doc() -> Path:
    doc = Document()
    doc.styles['Normal'].font.name = '맑은 고딕'
    doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')

    title = doc.add_paragraph()
    r = title.add_run('정보공개 청구 지원 시스템 보안성 검토서')
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor.from_string('0F172A')

    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate([
        ('대상 기관', '한국전력공사'),
        ('대상 시스템', '정보공개 청구 지원 시스템 Dashboard'),
        ('작성 기준일', '2026-09-03'),
        ('검토 등급', '초안 / 현장 정책 및 계정 운영 기준 반영 필요'),
    ]):
        doc_set_cell_text(meta.cell(i, 0), k, True, 'FFFFFF', 9)
        doc_set_cell_shading(meta.cell(i, 0), '0F7D80')
        doc_set_cell_text(meta.cell(i, 1), v, False, '111827', 9)

    def h(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(text)
        r.bold = True
        r.font.name = '맑은 고딕'
        r._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
        r.font.size = Pt(14)
        r.font.color.rgb = RGBColor.from_string('0F172A')

    def bullet(text):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.45)
        p.paragraph_format.first_line_indent = Cm(-0.25)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run('• ')
        r.font.color.rgb = RGBColor.from_string('0F7D80')
        r = p.add_run(text)
        r.font.name = '맑은 고딕'
        r._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
        r.font.size = Pt(10)

    h('1. 시스템 개요')
    for t in [
        '대시보드 서버는 기본 8000번 포트에서 실행되며, 담당자는 브라우저로 청구 접수, 판단 위저드, 통지 이력을 확인한다.',
        'AI API는 외부 주소를 입력받는 방식이 아니라 동일 장비의 8011번 포트에서 로컬 전용으로 구동된다.',
        '최종 통지 유형(공개/부분공개/비공개/정보부존재/진정질의/종결)은 AI가 자동 확정하지 않고, 반복 청구 판정을 제외한 모든 단계를 담당자가 근거 조문과 함께 확정한다.',
        '청구 원문, 판단 이력, 통지 결과는 로컬 SQLite DB에 저장된다.',
    ]:
        bullet(t)

    h('2. 보안상 강점')
    strengths = [
        ('외부 AI API 제거', '담당자가 외부 AI 서버 주소를 입력하지 않아도 되므로 오입력, 임의 서버 전송, API URL 노출 위험을 낮춘다.'),
        ('로컬 루프백 호출', 'AI API를 127.0.0.1/localhost 중심으로 제한하면 네트워크 외부 접근면을 줄일 수 있다.'),
        ('판단 위저드', '반복 청구를 제외한 모든 최종 판단은 담당자의 명시적 Y/N 확정을 거치며, 이미 완료된 청구는 재판단이 차단된다.'),
        ('이력 저장', '단계별 답변·근거 조문·담당자를 이력으로 남겨 사후 확인과 추적성을 제공한다.'),
    ]
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(tbl.rows[0].cells, ['항목', '보안 효과']):
        doc_set_cell_text(cell, text, True, 'FFFFFF', 9)
        doc_set_cell_shading(cell, '0F7D80')
    for k, v in strengths:
        row = tbl.add_row().cells
        doc_set_cell_text(row[0], k, True, '111827', 9)
        doc_set_cell_text(row[1], v, False, '111827', 9)

    h('3. 주요 위험 및 권고 통제')
    risks = [
        ('청구인 개인정보 노출', '높음', '청구 원문·연락처 열람 권한을 관리자 세션으로 제한, DB 파일 접근권한 최소화'),
        ('판단 오확정', '높음', '판단 위저드에 근거 조문을 항상 표시, 완료된 청구 재판단 차단(409 응답), 이력 상시 검토'),
        ('반복 청구 오탐/누락', '중간', '유사도 임계값을 현장 기준에 맞게 조정, 자동 판정 결과를 담당자가 상시 확인'),
        ('로컬 포트 외부 노출', '중간', 'AI API와 대시보드를 127.0.0.1 바인딩 또는 방화벽 제한, 필요 시 인증 추가'),
        ('SQLite DB 정보 노출', '중간', 'DB 파일 저장 위치 접근권한 제한, 백업/폐기 기준 마련'),
        ('EXE 변조/교체', '중간', '해시값 배포, 배포 폴더 쓰기 권한 제한'),
    ]
    tbl = doc.add_table(rows=1, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(tbl.rows[0].cells, ['위험', '등급', '권고 통제']):
        doc_set_cell_text(cell, text, True, 'FFFFFF', 8)
        doc_set_cell_shading(cell, '0F7D80')
    for risk, grade, control in risks:
        row = tbl.add_row().cells
        doc_set_cell_text(row[0], risk, True, '111827', 8)
        doc_set_cell_text(row[1], grade, True, 'B91C1C' if grade == '높음' else '92400E', 8)
        doc_set_cell_text(row[2], control, False, '111827', 8)

    h('4. 운영 전 확인 체크리스트')
    checks = [
        '대시보드 및 AI API 바인딩 주소가 외부 인터페이스로 열리지 않았는지 확인한다.',
        '관리자 계정 비밀번호 정책과 공유 여부를 점검한다.',
        '비공개 사유 키워드 목록과 통지서 템플릿 문구가 현장 기준과 최신 법령에 맞는지 확인한다.',
        '판단 이력 저장 위치와 보관 기간을 운영 정책에 맞춘다.',
        'EXE 배포본의 해시값을 기록한다.',
    ]
    for t in checks:
        bullet(t)

    h('5. 결론')
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(
        '현재 구조는 외부 AI API 의존을 줄이고 로컬 실행 중심으로 단순화했으며, 최종 판단의 대부분을 담당자 확정에 맡긴다는 점에서 '
        '오판단으로 인한 정보공개법 위반 위험을 낮춘다. 다만 청구인 개인정보를 다루는 시스템인 만큼 접근권한 통제와 이력 보존 정책이 '
        '함께 적용되어야 하며, 반복 청구 자동 판정 로직은 운영 데이터로 주기적으로 검증할 것을 권고한다.'
    )
    r.font.name = '맑은 고딕'
    r._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    r.font.size = Pt(10)

    OUT.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_PATH)
    return DOCX_PATH


if __name__ == '__main__':
    path = build_security_doc()
    print(path)
